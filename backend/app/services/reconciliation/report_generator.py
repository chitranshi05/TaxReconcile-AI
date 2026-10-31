from io import BytesIO
import math
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _transaction_row(transaction: dict[str, Any]) -> list[Any]:
    return [
        transaction.get("transaction_date", ""),
        transaction.get("description", ""),
        transaction.get("amount", ""),
        transaction.get("debit", ""),
        transaction.get("credit", ""),
    ]


def _pdf_text(value: Any) -> str:
    return escape(str(value))


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _xlsx_money(value: Any) -> Any:
    if not isinstance(value, str) or not value.strip():
        return value

    text = value.strip()
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1].strip()

    for symbol in ("₹", "$", "€", "£", ",", " "):
        text = text.replace(symbol, "")

    try:
        amount = float(text)
    except ValueError:
        return _xlsx_value(value)

    if not math.isfinite(amount):
        return _xlsx_value(value)
    return -abs(amount) if negative else amount


def _pdf_unmatched_table(
    transactions: list[dict[str, Any]],
    styles: Any,
) -> Table:
    rows = [["Date", "Description", "Amount", "Debit", "Credit"]]
    for transaction in transactions:
        rows.append(
            [
                _pdf_text(transaction.get("transaction_date", "")),
                Paragraph(
                    _pdf_text(transaction.get("description", "")),
                    styles["BodyText"],
                ),
                _pdf_text(transaction.get("amount", "")),
                _pdf_text(transaction.get("debit", "")),
                _pdf_text(transaction.get("credit", "")),
            ]
        )
    if len(rows) == 1:
        rows.append(["None", "", "", "", ""])
    table = Table(rows, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24476b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return table


def generate_reconciliation_xlsx(record: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Reconciliation", "Value"])
    summary_sheet.append(["Bank document", _xlsx_value(record.get("bank_filename", ""))])
    summary_sheet.append(["Ledger document", _xlsx_value(record.get("ledger_filename", ""))])
    summary_sheet.append(["Created at", _xlsx_value(str(record.get("created_at", "")))])
    for key, value in record.get("summary", {}).items():
        summary_sheet.append(
            [
                _xlsx_value(key.replace("_", " ").title()),
                _xlsx_money(value)
                if key in {
                    "total_bank_amount",
                    "total_ledger_amount",
                    "matched_amount",
                    "net_discrepancy",
                }
                else _xlsx_value(value) if isinstance(value, str) else value,
            ]
        )
    summary_sheet.freeze_panes = "A2"

    match_sheet = workbook.create_sheet("Matches")
    match_sheet.append(
        [
            "Status",
            "Match type",
            "Bank date",
            "Bank description",
            "Bank amount",
            "Ledger date",
            "Ledger description",
            "Ledger amount",
            "Amount variance",
            "Date difference (days)",
            "Confidence",
            "Notes",
        ]
    )
    for match in record.get("matches", []):
        bank = match.get("bank_transaction") or {}
        ledger = match.get("ledger_transaction") or {}
        match_sheet.append(
            [
                _xlsx_value(value) if isinstance(value, str) else value
                for value in [
                    match.get("status", ""),
                    match.get("match_type", ""),
                    bank.get("transaction_date", ""),
                    bank.get("description", ""),
                    _xlsx_money(
                        bank.get("amount", bank.get("debit", bank.get("credit", "")))
                    ),
                    ledger.get("transaction_date", ""),
                    ledger.get("description", ""),
                    _xlsx_money(
                        ledger.get("amount", ledger.get("debit", ledger.get("credit", "")))
                    ),
                    _xlsx_money(match.get("amount_variance", "")),
                    match.get("date_difference_days", ""),
                    match.get("confidence_score", ""),
                    match.get("notes", ""),
                ]
            ]
        )

    headers = ["Date", "Description", "Amount", "Debit", "Credit"]
    for sheet_name, key in (
        ("Bank only", "bank_only"),
        ("Ledger only", "ledger_only"),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for transaction in record.get(key, []):
            sheet.append(
                [
                    _xlsx_money(value)
                    if index in (2, 3, 4)
                    else _xlsx_value(value) if isinstance(value, str) else value
                    for index, value in enumerate(_transaction_row(transaction))
                ]
            )

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _xlsx_value(cell.value)
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column) + 2,
                48,
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_reconciliation_pdf(record: dict[str, Any]) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Reconciliation Report", styles["Title"]),
        Paragraph(
            f"Bank: {_pdf_text(record.get('bank_filename', ''))} | "
            f"Ledger: {_pdf_text(record.get('ledger_filename', ''))}",
            styles["Normal"],
        ),
        Spacer(1, 0.15 * inch),
    ]

    summary = record.get("summary", {})
    summary_data = [["Metric", "Value"]] + [
        [key.replace("_", " ").title(), str(value)]
        for key, value in summary.items()
    ]
    summary_table = Table(summary_data, repeatRows=1)
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24476b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elements.extend([summary_table, Spacer(1, 0.2 * inch)])

    match_data = [
        [
            "Status",
            "Bank date / description / amount",
            "Ledger date / description / amount",
            "Variance",
            "Date diff.",
            "Confidence",
        ]
    ]
    for match in record.get("matches", []):
        bank = match.get("bank_transaction") or {}
        ledger = match.get("ledger_transaction") or {}
        match_data.append(
            [
                _pdf_text(match.get("status", "")),
                Paragraph(
                    "<br/>".join(
                        _pdf_text(value)
                        for value in (
                            bank.get("transaction_date", ""),
                            bank.get("description", ""),
                            bank.get("amount", bank.get("debit", bank.get("credit", ""))),
                        )
                    ),
                    styles["BodyText"],
                ),
                Paragraph(
                    "<br/>".join(
                        _pdf_text(value)
                        for value in (
                            ledger.get("transaction_date", ""),
                            ledger.get("description", ""),
                            ledger.get("amount", ledger.get("debit", ledger.get("credit", ""))),
                        )
                    ),
                    styles["BodyText"],
                ),
                _pdf_text(match.get("amount_variance", "")),
                _pdf_text(match.get("date_difference_days", "")),
                _pdf_text(match.get("confidence_score", "")),
            ]
        )
    if len(match_data) == 1:
        match_data.append(["No matches", "", "", "", "", ""])
    match_table = Table(match_data, repeatRows=1)
    match_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24476b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.extend([Paragraph("Matches", styles["Heading2"]), match_table])
    for heading, key in (
        ("Bank-only transactions", "bank_only"),
        ("Ledger-only transactions", "ledger_only"),
    ):
        elements.extend(
            [
                Paragraph(heading, styles["Heading2"]),
                _pdf_unmatched_table(record.get(key, []), styles),
            ]
        )
    document.build(elements)
    return output.getvalue()
